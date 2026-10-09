"""Excel export of a screening run (FR-X7): ranking, comparison matrix and evidence, one sheet each.

Why this file exists:
- Recruiters often share results in Excel. This builds a tidy workbook with
  bold headers, sensible column widths and a frozen top row.
- The Comparison sheet colours each cell like the app (green met, amber
  partial, red missing) AND writes the symbol, so colour is never the only signal.
- It shows the AI's band next to the recruiter's decision, never in place of it.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from core.report import candidate_name, decision_for, missing_info_names, screened_candidates
from core.schemas import RunResult

STATUS_FILLS: dict[str, tuple[str, str]] = {  # status -> (cell text, background colour)
    "met": ("✔ met", "E5F4EC"),
    "partial": ("◐ partial", "FFF4D6"),
    "missing": ("✖ missing", "FDE7E5"),
}
HEADER_FILL = PatternFill("solid", fgColor="E6F3F2")
MAX_COLUMN_WIDTH = 60


def style_sheet(sheet: Worksheet) -> None:
    """Bold headers, frozen top row, and column widths that fit the text (up to a limit)."""
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="0A5E5D")
        cell.fill = HEADER_FILL
    sheet.freeze_panes = "A2"
    for column_cells in sheet.columns:
        longest = max(len(str(cell.value or "")) for cell in column_cells)
        letter = get_column_letter(column_cells[0].column)
        sheet.column_dimensions[letter].width = min(max(longest + 2, 8), MAX_COLUMN_WIDTH)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def add_ranking_sheet(workbook: Workbook, run: RunResult, decisions: dict[str, dict]) -> None:
    sheet = workbook.active
    sheet.title = "Ranking"
    sheet.append(["Rank", "Name", "Score", "AI band (recommendation)", "Must-have %", "Nice-to-have %",
                  "Missing information", "Flags", "Recruiter decision"])
    for rank, candidate in enumerate(screened_candidates(run), start=1):
        score = candidate.score
        sheet.append([
            rank,
            candidate_name(candidate),
            score.overall,
            score.band,
            score.must_have_pct,
            score.nice_to_have_pct,
            ", ".join(missing_info_names(candidate)),
            ", ".join(flag.code for flag in candidate.flags),
            decision_for(candidate, decisions),
        ])
    style_sheet(sheet)


def add_comparison_sheet(workbook: Workbook, run: RunResult) -> None:
    sheet = workbook.create_sheet("Comparison")
    requirements = run.requirements.must_have + run.requirements.nice_to_have
    sheet.append(["Candidate"] + [f"{requirement.id} {requirement.label}" for requirement in requirements])
    for candidate in screened_candidates(run):
        statuses = {result.requirement_id: result.status for result in candidate.assessment.results}
        sheet.append([candidate_name(candidate)] + [STATUS_FILLS[statuses.get(r.id, "missing")][0] for r in requirements])
        for position, requirement in enumerate(requirements, start=2):
            colour = STATUS_FILLS[statuses.get(requirement.id, "missing")][1]
            sheet.cell(row=sheet.max_row, column=position).fill = PatternFill("solid", fgColor=colour)
    style_sheet(sheet)
    for position in range(2, len(requirements) + 2):  # keep the matrix compact
        sheet.column_dimensions[get_column_letter(position)].width = 16


def add_evidence_sheet(workbook: Workbook, run: RunResult) -> None:
    sheet = workbook.create_sheet("Evidence")
    sheet.append(["Candidate", "Requirement", "Status", "Quote from the CV", "Verified", "Guidelines"])
    labels = {r.id: f"{r.id} {r.label}" for r in run.requirements.must_have + run.requirements.nice_to_have}
    for candidate in screened_candidates(run):
        for result in candidate.assessment.results:
            sheet.append([
                candidate_name(candidate),
                labels.get(result.requirement_id, result.requirement_id),
                result.status,
                result.evidence or "",
                "yes" if result.verified else ("" if result.status == "missing" else "no"),
                ", ".join(result.guideline_refs),
            ])
    style_sheet(sheet)


def build_workbook(run: RunResult, decisions: dict[str, dict]) -> bytes:
    """The .xlsx file as bytes, ready for a download button."""
    workbook = Workbook()
    add_ranking_sheet(workbook, run, decisions)
    add_comparison_sheet(workbook, run)
    add_evidence_sheet(workbook, run)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
