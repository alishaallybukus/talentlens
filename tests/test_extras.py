"""Tests for the extras (Phase 8): Excel export, missing-information emails, and the Ask-the-CVs assistant.

Everything uses the FakeLLM, so no real AI is called.
"""

from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook

from core.agents import run_missing_info_email
from core.assistant import MAX_TOOL_CALLS, NOT_FOUND_ANSWER, CVTools, answer_question
from core.excel import build_workbook
from core.rag import build_index
from core.schemas import Citation, JobRequirements
from core.supervisor import create_run, run_screening
from eval.run_eval import load_cvs
from tests.conftest import SAMPLE_DIR
from tests.fake_llm import SAMPLE_REQUIREMENTS, FakeLLM


@pytest.fixture(scope="module")
def screened():
    """A run of all 7 CVs with the fake model, plus each CV's cleaned text by candidate id."""
    requirements = JobRequirements.model_validate(SAMPLE_REQUIREMENTS)
    index = build_index((SAMPLE_DIR / "hiring_guidelines.md").read_text(encoding="utf-8"))
    documents = load_cvs(None)
    run = create_run(requirements, "fake", "fake-model")
    run_screening(run, documents, FakeLLM(), index, today=date(2026, 10, 9))
    texts = {intake.content_hash[:12]: intake.clean_text for intake in documents}
    return run, texts


def candidate(run, first_name: str):
    return next(c for c in run.candidates if c.profile.name.startswith(first_name))


# --- 8a Excel ---------------------------------------------------------------------


def test_excel_has_three_sheets_with_coloured_matrix(screened) -> None:
    run, _texts = screened
    sarah = candidate(run, "Sarah")
    workbook = load_workbook(BytesIO(build_workbook(run, {sarah.candidate_id: {"decision": "Shortlist"}})))
    assert workbook.sheetnames == ["Ranking", "Comparison", "Evidence"]

    ranking = workbook["Ranking"]
    assert ranking["A1"].font.bold and ranking.freeze_panes == "A2"
    assert ranking.max_row == 8  # header + 7 candidates
    decisions = {row[1]: row[8] for row in ranking.iter_rows(min_row=2, values_only=True)}
    assert decisions[sarah.profile.name] == "Shortlist"
    assert set(decisions.values()) == {"Shortlist", "Pending"}

    comparison = workbook["Comparison"]
    cell = comparison["B2"]
    assert cell.value in ("✔ met", "◐ partial", "✖ missing")  # a symbol, not colour alone
    assert cell.fill.fgColor.rgb.endswith(("E5F4EC", "FFF4D6", "FDE7E5"))

    evidence = workbook["Evidence"]
    assert [c.value for c in evidence[1]] == ["Candidate", "Requirement", "Status", "Quote from the CV", "Verified", "Guidelines"]


# --- 8b email drafts ----------------------------------------------------------------


def test_kevin_email_asks_for_exactly_his_missing_items() -> None:
    fake = FakeLLM()
    missing = ["salary expectation", "notice period or availability", "right to work in Mauritius"]
    step = run_missing_info_email(fake, "Marketing Executive", "Corallia Living Ltd", "Kevin Ramdin", missing, "Alisha")
    for item in missing:
        assert item in step.result.body
    sent = fake.everything_sent()
    assert "Kevin Ramdin" in sent
    assert "<cv>" not in fake.calls[0]["prompt"]  # no CV text is sent for an email
    assert "\n- " in step.result.body  # the bullet list survives the bias filter


def test_email_bias_filter_removes_protected_sentences() -> None:
    def chatty_email(prompt):
        return {"subject": "Details", "body": "Thank you for applying.\nAre you married? We'd like to know.\n- notice period"}

    step = run_missing_info_email(FakeLLM({"email": chatty_email}), "Job", "Co", "Name", ["notice period"], "R")
    assert "married" not in step.result.body
    assert "- notice period" in step.result.body


# --- 8c Ask the CVs -----------------------------------------------------------------


def test_assistant_calls_a_tool_then_answers_with_a_verified_quote(screened) -> None:
    run, texts = screened
    tools = CVTools(run.candidates, texts)
    answer = answer_question(FakeLLM(), "TikTok", tools)
    assert [call.tool for call in answer.tool_calls] == ["search_cvs"]
    assert answer.tool_calls[0].label() == "🔧 search_cvs('TikTok')"
    assert answer.citations and all(citation.verified for citation in answer.citations)
    assert "TikTok" in answer.citations[0].quote


def test_invented_quote_is_not_verified(screened) -> None:
    run, texts = screened
    tools = CVTools(run.candidates, texts)
    sarah = candidate(run, "Sarah")
    assert not tools.verify(Citation(candidate=sarah.profile.name, quote="Ran a team of 50 engineers at NASA"))


def test_tool_calls_are_limited(screened) -> None:
    run, texts = screened

    def always_searching(prompt):
        if "used all your tool calls" in prompt:
            return {"action": "answer", "answer": "Done searching.", "citations": []}
        return {"action": "call_tool", "tool": "search_cvs", "arguments": {"query": "marketing"}}

    answer = answer_question(FakeLLM({"assistant": always_searching}), "Anything?", CVTools(run.candidates, texts))
    assert len(answer.tool_calls) == MAX_TOOL_CALLS
    assert answer.answer == "Done searching."


def test_unanswerable_question(screened) -> None:
    run, texts = screened
    answer = answer_question(FakeLLM(), "zzqx", CVTools(run.candidates, texts))
    assert answer.answer == NOT_FOUND_ANSWER


def test_tools_find_candidates_and_handle_bad_input(screened) -> None:
    run, texts = screened
    tools = CVTools(run.candidates, texts)
    assert "Kevin Ramdin" in tools.run("get_candidate", {"name": "kevin"})
    assert "No candidate called" in tools.run("get_candidate", {"name": "Nobody"})
    assert "There is no tool" in tools.run("delete_everything", {})
    assert "Wrong arguments" in tools.run("search_cvs", {"nonsense": 1})
    listed = tools.run("list_candidates", {})
    assert listed.count("\n") == 6  # 7 candidates


def test_ryan_hidden_text_is_not_searchable(screened) -> None:
    """The tools only see the cleaned CV text, so quarantined injection lines can't be found or quoted."""
    run, texts = screened
    ryan = candidate(run, "Ryan")
    tools = CVTools(run.candidates, texts)
    for line in ryan.quarantined_text:
        assert line not in tools.search_cvs(line)
